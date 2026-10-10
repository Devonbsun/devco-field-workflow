"""Close status and explicit billing stay independent through UI and exports."""
import ast
import contextlib
import io
import json
import re
import tempfile
import threading
import unittest
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import urlopen
from unittest.mock import Mock
from openpyxl import load_workbook

import devco_app as app
from invoice_total import invoice_summary
from job_records import collect, _parse_record
from project_directory import billing_workbook
from voice_workflow import CLOSE_LABELS


class ManualBilling(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        app.JOBS=self.root/'JOBS';app.STATE=self.root/'state.json';app.active_job='TEST';app.active_ju='101'
        self.folder=app.JOBS/'TEST'/'3_JU_FILES'/'101 - Test'
        (self.folder/'photos').mkdir(parents=True)
        (self.folder/'transfer_info.txt').write_text('JU Record: 101\nAddress: Test Street\nLatitude: 41.6\nLongitude: -93.6\n')
        for name in ['one.jpg','two.jpg']:(self.folder/'photos'/name).write_bytes(b'photo')
        self.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.H)
        threading.Thread(target=self.server.serve_forever,daemon=True).start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.tmp.cleanup()
    def close(self,status,codes):
        app.active_ju='101'
        data=[('job','TEST'),('ju','101'),('close',status),('note','Access blocked')]
        for code,qty in codes:data.extend([('code',code),('qty_'+code,str(qty))])
        with urlopen(self.base+'/finish',data=urlencode(data).encode()) as r:r.read()
        return _parse_record(self.folder/'BILLING_AND_NOTES.txt')
    def test_unable_closes_ju_in_maps_sheets_and_daily_exports(self):
        rec=self.close('UNABLE TO COMPLETE',[])
        self.assertEqual(rec['status'],'UNABLE TO COMPLETE');self.assertTrue(rec['completed_at'])
        self.assertEqual(app._ju_state(self.folder),'COMPLETE');self.assertTrue(app.map_record('TEST','101')['done'])
        row=collect(app.JOBS/'TEST')[0];self.assertEqual(row['State'],'COMPLETE')
        wb=load_workbook(app.JOBS/'TEST'/'1_JOB_WORKFLOW'/'TEST_COMPLETED.xlsx')
        self.assertEqual(wb.active.cell(2,1).value,'101');self.assertEqual(wb.active.cell(2,7).value,'UNABLE TO COMPLETE');wb.close()
        from daily_export import records as completed_records
        # Daily export reads the shared COMPLETE rule, preserving the closing task.
        rows,_=completed_records(self.root,'TEST')
        self.assertTrue(any(r['ju']=='101' and r['close_code']=='UNABLE TO COMPLETE' for r in rows))
    def test_close_codes_never_add_or_replace_billing(self):
        statuses=['TRANSFER ALREADY COMPLETED','NO SERVICES ON POLE','NO IDENTIFIABLE WINDSTREAM LINE ON POLE','ADSS','NO WINDSTREAM VIOLATION','UNABLE TO COMPLETE']
        for close in statuses:
            with self.subTest(close=close):
                rec=self.close(close,[]);self.assertEqual(rec['billing'],[])
                self.assertEqual(invoice_summary(app.JOBS/'TEST')['trip_quantity'],0)
                self.assertEqual(float(invoice_summary(app.JOBS/'TEST')['total']),0)
                selected=[('TRIP CHARGE','2'),('WC1','1')]
                rec=self.close(close,selected);self.assertEqual(rec['billing'],selected)
                invoice=invoice_summary(app.JOBS/'TEST');self.assertEqual(invoice['trip_quantity'],2)
                self.assertEqual(float(invoice['trip']),80);self.assertEqual(float(invoice['total']),119.18)
    def test_old_saved_trip_lines_stay_billable_without_inference(self):
        file=self.folder/'BILLING_AND_NOTES.txt'
        for billing,quantity in [('No billing codes entered',0),('TRIP CHARGE x1',1)]:
            file.write_text('STATUS: ADSS\nBILLING:\n'+billing+'\nNOTES:\nSurveyed\n')
            before=file.read_bytes();summary=invoice_summary(app.JOBS/'TEST')
            self.assertEqual(summary['trip_quantity'],quantity)
            output=self.root/'billing.xlsx';billing_workbook(app.JOBS/'TEST',collect(app.JOBS/'TEST'),output)
            wb=load_workbook(output);ws=wb['Billing']
            codes=[ws.cell(r,4).value for r in range(2,ws.max_row+1)]
            self.assertEqual(codes.count('TRIP CHARGE'),quantity)
            self.assertEqual(file.read_bytes(),before);wb.close()
    def test_trip_appears_only_as_selectable_billing(self):
        page=app.billing_page()
        self.assertIn('name="code" value="TRIP CHARGE"',page)
        self.assertIn('name="qty_TRIP CHARGE" value="1"',page)
        self.assertEqual(app.BILLING_CODES.count('TRIP CHARGE'),1)
        for label in re.findall(r'<button[^>]*name="close"[^>]*>(.*?)</button>',page):self.assertNotIn('Trip',label)
        self.assertTrue(all('Trip' not in label for label in CLOSE_LABELS.values()))
    def test_terminal_choices_use_manual_trip_code(self):
        source=Path(__file__).with_name('field_workflow.py').read_text()
        tree=ast.parse(source);fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='billing_prompt')
        save=Mock();scope={'MASTER_LOG':self.root/'production.csv','re':re,'MASTER':self.folder.parent,'JOB':app.JOBS/'TEST','CODES':{'4':'TRIP CHARGE'},'save_master_record':save,'sync_job':Mock(return_value={'COMPLETE':1,'TOTAL':1,'PENDING':0}),'navigate_next':Mock()}
        exec(compile(ast.Module(body=[fn],type_ignores=[]),'<terminal closeout>','exec'),scope)
        for choices,expected in [(['6','','Surveyed'],[]),(['8','4','1','','Access blocked'],[('TRIP CHARGE','1')])]:
            answers=iter(choices)
            with contextlib.redirect_stdout(io.StringIO()):scope['billing_prompt']({'ju':'101','address':'Test Street'},list((self.folder/'photos').glob('*')),lambda _:next(answers))
            self.assertEqual(save.call_args.args[2],expected)


if __name__=='__main__':unittest.main()
