"""Exercise actual saves, conflict protection, exports, and photo routes on fake jobs."""
import json
import threading
import tempfile
import unittest
import urllib.request
import urllib.parse
from urllib.error import HTTPError
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import devco_app as app
from record_editor import record_data
from job_records import collect, _parse_record

class RecordEditor(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        app.JOBS=Path(self.tmp.name)/'JOBS';app.STATE=Path(self.tmp.name)/'state.json'
        app.active_job='TEST';app.active_ju='101'
        self.folder=app.JOBS/'TEST'/'3_JU_FILES'/'101 - Test'
        self.folder.mkdir(parents=True)
        (self.folder/'transfer_info.txt').write_text('JU Record: 101\nAddress: Test Street\nLatitude: 41.6\nLongitude: -93.6\n')
        (self.folder/'photos').mkdir()
        for name in ['one.jpg','two.jpg']:(self.folder/'photos'/name).write_bytes(b'test-photo')
        app.save_closeout('TEST','101','FIBER TRANSFER COMPLETED',[('WC1','2')],'Original note')
        self.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.H)
        threading.Thread(target=self.server.serve_forever,daemon=True).start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.tmp.cleanup()
    def get(self,path):return urllib.request.urlopen(self.base+path).read()
    def payload(self):
        d=record_data(vars(app),'TEST','101')
        return {k:d[k] for k in ['job','ju','revision','note','billing']}
    def save(self,data):
        req=urllib.request.Request(self.base+'/record-save',data=json.dumps(data).encode(),headers={'Content-Type':'application/json'})
        return json.loads(urllib.request.urlopen(req).read())
    def test_edit_saves_existing_ju_after_active_ju_changes(self):
        before=collect(app.JOBS/'TEST')[0]['Completed At']
        data=self.payload();data.update(note='Corrected note\nWC1 x999',billing=[['WC1','3'],['PE1-3G','2']])
        app.active_ju='999'
        result=self.save(data)
        self.assertTrue(result['ok']);self.assertEqual(app.active_ju,'999')
        record=record_data(vars(app),'TEST','101')
        self.assertEqual(record['note'],data['note']);self.assertEqual(record['status'],'FIBER TRANSFER COMPLETED')
        self.assertEqual(record['billing'],[('WC1','3'),('PE1-3G','2')])
        wb=load_workbook(app.JOBS/'TEST'/'1_JOB_WORKFLOW'/'TEST_MASTER.xlsx')
        self.assertEqual(wb.active.cell(2,9).value,'WC1 x3; PE1-3G x2')
        self.assertEqual(wb.active.cell(2,10).value,data['note']);wb.close()
        self.assertEqual(collect(app.JOBS/'TEST')[0]['Completed At'],before)
        self.assertEqual(len(list((self.folder/'RECORD_HISTORY').glob('*.json'))),1)
    def test_stale_save_cannot_overwrite(self):
        old=self.payload();fresh=self.payload();fresh['note']='Saved first';self.save(fresh)
        old['note']='Stale overwrite'
        with self.assertRaises(HTTPError) as caught:self.save(old)
        self.assertEqual(caught.exception.code,409)
        self.assertEqual(record_data(vars(app),'TEST','101')['note'],'Saved first')
    def test_invalid_billing_does_not_change_record(self):
        before=(self.folder/'BILLING_AND_NOTES.txt').read_bytes()
        for billing in [[['WC1','0']],[['WC1','1.5']],[['BAD','1']],[['WC1','1'],['WC1','2']],[]]:
            data=self.payload();data['billing']=billing
            with self.assertRaises(HTTPError) as caught:self.save(data)
            self.assertEqual(caught.exception.code,400)
            self.assertEqual(before,(self.folder/'BILLING_AND_NOTES.txt').read_bytes())
    def test_draft_conflict_and_sync_warning(self):
        stale=self.payload();app.save_note(self.folder,'New field draft')
        with self.assertRaises(HTTPError) as caught:self.save(stale)
        self.assertEqual(caught.exception.code,409)
        data=self.payload();data['note']='Saved despite export failure'
        with patch.object(app,'sync_job',side_effect=OSError('disk')):
            result=self.save(data)
        self.assertIn('Spreadsheet refresh failed',result['warning'])
        self.assertEqual(record_data(vars(app),'TEST','101')['note'],data['note'])
    def test_photo_bytes_and_traversal_rejected(self):
        items=json.loads(self.get('/photos?job=TEST&ju=101'))
        self.assertEqual(len(items),2)
        self.assertEqual(self.get(items[0]['url']),b'test-photo')
        for name in ['../transfer_info.txt','/etc/passwd','one.jpg/../../transfer_info.txt']:
            with self.assertRaises(HTTPError) as caught:self.get('/photo?'+urllib.parse.urlencode(dict(job='TEST',ju='101',name=name)))
            self.assertEqual(caught.exception.code,404)
        external=Path(self.tmp.name)/'external.jpg';external.write_bytes(b'private')
        (self.folder/'photos'/'link.jpg').symlink_to(external)
        with self.assertRaises(HTTPError):self.get('/photo?job=TEST&ju=101&name=link.jpg')
        self.assertEqual(len(json.loads(self.get('/photos?job=TEST&ju=101'))),2)
    def test_no_gps_record_is_listed_and_can_be_edited(self):
        folder=self.folder.parent/'202 - No GPS';folder.mkdir()
        (folder/'transfer_info.txt').write_text('JU Record: 202\nAddress: No GPS <address>\n')
        self.assertIn('JU 202',self.get('/records?job=TEST').decode())
        self.assertIn('No GPS &lt;address&gt;',self.get('/record?job=TEST&ju=202').decode())
        data=record_data(vars(app),'TEST','202');data['billing']=[['TRIP CHARGE','1']];data['note']='Survey pending'
        result=self.save(data);self.assertTrue(result['ok'])
        self.assertEqual(app._ju_state(folder),'NOT COMPLETED')
    def test_native_camera_fallback_targets_verified_component(self):
        with patch.object(app.subprocess,'run') as run:
            run.return_value.returncode=0;run.return_value.stdout='';run.return_value.stderr=''
            self.get('/camera')
            args=run.call_args.args[0]
            self.assertIn('com.solocator/com.solocator.splash.SplashActivity',args)
            self.assertEqual(args[args.index('--user')+1],'0')
    def test_record_content_is_escaped(self):
        data=self.payload();data['note']='</script><script>alert(1)</script>';self.save(data)
        page=self.get('/record?job=TEST&ju=101').decode()
        self.assertNotIn(data['note'],page)
        self.assertIn('<\\/script>',page)

if __name__=='__main__':unittest.main()
