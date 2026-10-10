import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image
from openpyxl import load_workbook
import daily_export as d


class DailyExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.job = 'TEST'

    def tearDown(self):
        self.tmp.cleanup()

    def pole(self, ju, dates=('2026-10-03',), completed='', status='NO SERVICES ON POLE'):
        folder = self.root / 'JOBS/TEST/3_JU_FILES' / ju
        photos = folder / 'photos'; photos.mkdir(parents=True)
        (folder / 'transfer_info.txt').write_text(f'JU Record: {ju}\nAddress: =danger\nLatitude: 41.0\nLongitude: -93.0\n')
        (folder / 'BILLING_AND_NOTES.txt').write_text(f'COMPLETED_AT: {completed}\nSTATUS: {status}\nBILLING:\nTRIP CHARGE x1\nNOTES:\nSaved note\n')
        (folder / 'NOTE_DRAFT.json').write_text('{"note":"SECRET DRAFT"}')
        (folder / 'VOICE_HISTORY').mkdir(); (folder / 'VOICE_HISTORY/private.txt').write_text('SECRET VOICE')
        for n, day in enumerate(dates):
            Image.new('RGB', (60, 60), 'red').save(photos / f'Solocator-{day} 12-00-0{n}.jpg')
        return folder

    def build(self, limit=d.LIMIT):
        rows, _ = d.records(self.root, self.job)
        value = {'id': 'test-export', 'status': 'building', 'count': len(rows), 'progress': 0}
        d.build(self.root, self.job, rows, value, limit)
        return json.loads(d.status_path(self.root, self.job).read_text())

    def test_date_precedence_and_completed_only(self):
        self.pole('1', completed='2026-10-04T02:00:00Z') # still October 3 Central
        self.pole('2', dates=('2026-09-29', '2026-10-05'))
        self.pole('3', status='PENDING')
        self.pole('4', status='FIBER TRANSFER COMPLETED') # insufficient photos
        rows, _ = d.records(self.root, self.job)
        self.assertEqual([(r['ju'], r['day']) for r in rows], [('1','2026-10-03'), ('2','2026-10-05')])
        self.assertEqual(rows[0]['date_source'], 'Saved completion timestamp')
        self.assertEqual(rows[1]['date_source'], 'Solocator photo filename timestamp')
        self.assertEqual(d.completion_day('not a date'), d.UNKNOWN)

    def test_export_coverage_layout_and_safe_excel(self):
        self.pole('1'); self.pole('2'); self.pole('3', dates=('2026-10-04',))
        result = self.build(limit=20000)
        self.assertEqual(result['status'], 'ready', result)
        self.assertEqual([(day['day'],day['count']) for day in result['days']], [('2026-10-03',2),('2026-10-04',1)])
        all_names = []
        for day in result['days']:
            for part in day['parts']:
                self.assertLess(part['bytes'], 20000)
                path = d.packet_file(self.root, self.job, result['id'], part['name'])
                with zipfile.ZipFile(path) as z:
                    self.assertIsNone(z.testzip()); all_names.extend(z.namelist())
                    self.assertTrue(all(n.startswith(day['day']+'/') for n in z.namelist()))
                    for name in z.namelist():
                        self.assertNotIn('NOTE_DRAFT',name); self.assertNotIn('VOICE_HISTORY',name)
                        if name.endswith('Completed_JUs.xlsx'):
                            ws=load_workbook(io.BytesIO(z.read(name))).active
                            self.assertEqual(ws['C2'].value, '=danger'); self.assertEqual(ws['C2'].data_type,'s')
                            self.assertIsNotNone(ws['D2'].value.year)
        self.assertEqual(len(all_names),len(set(all_names)))
        self.assertEqual(sum(n.endswith('JU_REPORT.txt') for n in all_names),3)
        self.assertEqual(sum(n.endswith('.jpg') for n in all_names),3)
        self.assertEqual(sum('/3_Billing_Only/' in n for n in all_names),3)
        self.assertTrue((d.destination(self.root,self.job)/'2026-10-03/1_Spreadsheets/Completed_JUs.xlsx').exists())
        with self.assertRaises(ValueError): d.packet_file(self.root,self.job,'stale',result['days'][0]['parts'][0]['name'])
        with self.assertRaises(ValueError): d.packet_file(self.root,self.job,result['id'],'../../secret')

    def test_source_changes_reject_packet_and_repeat_archives_previous(self):
        f=self.pole('1'); first=self.build(); self.assertEqual(first['status'],'ready')
        original=d.packets.compact_photo
        def changed(root, source):
            result=original(root, source)
            with (f/'BILLING_AND_NOTES.txt').open('a') as stream: stream.write('changed\n')
            return result
        with patch.object(d.packets,'compact_photo',changed): result=self.build()
        self.assertEqual(result['status'],'failed')
        self.assertIn('changed',result['error'])
        self.assertTrue(d.destination(self.root,self.job).exists())
        self.assertEqual(self.build()['status'],'ready')
        self.assertTrue((self.root/'TRACKING/daily-history').exists())

    def test_original_photo_exif_wins_over_filename(self):
        f=self.pole('1')/'photos/Solocator-2026-10-03 12-00-00.jpg'
        exif=Image.Exif(); exif[36867]='2026:10:04 23:15:00'
        Image.new('RGB',(30,30)).save(f,exif=exif)
        self.assertEqual(d.photo_date(f), ('2026-10-04','Original photo timestamp'))


if __name__ == '__main__': unittest.main()
