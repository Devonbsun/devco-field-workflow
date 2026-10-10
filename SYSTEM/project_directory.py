"""Publish phone project folders in place from authoritative records.

Never create dated/final/(1) work packets. Only backups are versioned.
The 212 legacy packet does not have reliable standardized completion states.
"""
from pathlib import Path
import csv, fcntl, hashlib, io, json, os, shutil, threading, time, zipfile
from datetime import datetime
from tempfile import TemporaryDirectory
from openpyxl import Workbook
from job_records import collect, _write
from record_editor import RECORD_LOCK

REPO = Path.home() / "DEVCO_FIELD"
DIRECTORY = Path("/storage/emulated/0/DEVCO/Pole_Transfers")
BACKUPS = DIRECTORY.parent / "Backups"
STATE = REPO / ".project_directory_state.json"
PART_LIMIT = 95_000_000
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".gif"}
VIDEO_TYPES = {".mp4", ".mov", ".m4v", ".webm", ".3gp"}

def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(4 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def atomic_text(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name("." + path.name + ".tmp")
    tmp.write_text(text)
    tmp.replace(path)

def remember(path):
    """Preserve overwritten/deleted published contents by hash."""
    digest = sha(path)
    dest = BACKUPS / "Updated_Files" / "objects" / digest[:2] / digest
    if not dest.exists():
        dest.parent.mkdir(parents=True, exist_ok=True)
        tmp = dest.with_suffix(".tmp")
        shutil.copy2(path, tmp)
        if sha(tmp) != digest:
            raise RuntimeError("Changed-file backup failed")
        tmp.replace(dest)
    elif sha(dest) != digest:
        raise RuntimeError("Existing changed-file backup is corrupt")
    with (BACKUPS / "Updated_Files" / "index.jsonl").open("a") as f:
        f.write(json.dumps({"path": str(path), "sha256": digest, "time": datetime.now().isoformat()}) + "\n")

def publish_file(source, target):
    target.parent.mkdir(parents=True, exist_ok=True)
    desired = sha(source)
    if target.exists() and sha(target) == desired:
        return desired
    if target.exists():
        remember(target)
    temp = target.with_name("." + target.name + ".tmp")
    shutil.copy2(source, temp)
    if sha(temp) != desired:
        raise RuntimeError("Copy verification failed: " + str(target))
    temp.replace(target)
    return desired

def source_files(job):
    root = REPO / "JOBS" / job
    legacy = root / "JOB_PACKET_CURRENT"
    folders = [root / s for s in ("0_ORIGINAL", "1_JOB_WORKFLOW", "2_ROUTE")]
    folders.append(legacy if legacy.exists() else root / "3_JU_FILES")
    result = []
    for folder in folders:
        if folder.exists():
            result.extend(p for p in folder.rglob("*") if p.is_file()
                          and not p.is_symlink() and "__pycache__" not in p.parts)
    if (root / "JOB_PACKET_CURRENT_SUMMARY.json").exists():
        result.append(root / "JOB_PACKET_CURRENT_SUMMARY.json")
    return sorted(result)

def fingerprint(job):
    return hashlib.sha256(json.dumps([(str(p), p.stat().st_size, p.stat().st_mtime_ns)
                         for p in source_files(job)]).encode()).hexdigest()

def billing_workbook(source, rows, target):
    from invoice_total import RATES, invoice_summary
    from job_records import _parse_record
    wb = Workbook()
    ws = wb.active
    ws.title = "Billing"
    ws.append(["JU", "Address", "Closing task", "Billing code", "Quantity", "Unit rate",
               "Line total", "Notes", "Photos", "Completed at"])
    for row in rows:
        if row["State"] != "COMPLETE":
            continue
        folder = next((p.parent for p in (source / "3_JU_FILES").glob(str(row["JU"]) + " - */transfer_info.txt")), None)
        if folder is None:
            raise RuntimeError("Billing JU folder missing")
        codes = list(_parse_record(folder / "BILLING_AND_NOTES.txt")["billing"])
        if not codes:
            codes = [("MISSING BILLING CODE", "0")]
        for code, quantity in codes:
            qty = int(quantity)
            rate = RATES.get(code.strip())
            ws.append([row["JU"], row["Address"], row["Close Code"], code, qty, rate,
                       round(qty * rate, 2) if rate is not None else None,
                       row["Notes"], row["Photo Count"], row["Completed At"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJ", [15, 50, 42, 26, 12, 14, 14, 70, 10, 25]):
        ws.column_dimensions[col].width = width
    for values in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in values:
            cell.number_format = '"$"#,##0.00'
    rates = wb.create_sheet("Rates")
    rates.append(["Code", "Rate used by field app"])
    for code, rate in sorted(RATES.items()):
        rates.append([code, rate])
    total = invoice_summary(source)
    summary = wb.create_sheet("Summary")
    for row in [["Project", source.name], ["Priced total", float(total["total"])],
                ["Work", float(total["work"])], ["Trip charges", float(total["trip"])],
                ["Trip quantity", total["trip_quantity"]], ["Billed JUs", total["billed_jus"]],
                ["Unpriced codes", json.dumps(total["unpriced"])],
                ["Source", "Current saved JU records; same rules and rates as field app"]]:
        summary.append(row)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 80
    wb.save(target)

def legacy_index(source, target):
    rows = json.loads((source / "JOB_PACKET_CURRENT_SUMMARY.json").read_text())
    if len(rows) != len({str(r["ju"]) for r in rows}):
        raise RuntimeError("Duplicate legacy JU identifiers")
    wb = Workbook()
    ws = wb.active
    ws.title = "JU Index"
    ws.append(["JU", "Address", "Recorded field data", "Completion status", "Billing as recorded", "Photos"])
    for row in rows:
        ws.append([row["ju"], row["address"], "Yes" if row["has_work_data"] else "No",
                   "Unverified in legacy record",
                   "; ".join(str(c) + " x" + str(q) for c, q in row["billing"].items()), row["photos"]])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEF", [16, 55, 24, 35, 55, 12]):
        ws.column_dimensions[col].width = width
    wb.save(target)
    return rows

def build_mapping(job, stage):
    source = REPO / "JOBS" / job
    mapping = {}
    sheet = stage / "Spreadsheets"
    sheet.mkdir()
    for section in ("0_ORIGINAL", "2_ROUTE"):
        for p in sorted((source / section).glob("*.xlsx")):
            mapping["Spreadsheets/" + p.name] = p
    legacy = (source / "JOB_PACKET_CURRENT").exists()
    counts = {}
    if legacy:
        records = source / "JOB_PACKET_CURRENT"
        for report in sorted((source / "1_JOB_WORKFLOW").glob("*.xlsx")):
            mapping["Spreadsheets/Saved_Reports/" + report.name] = report
        rows = legacy_index(source, sheet / (job + "_JU_INDEX.xlsx"))
        counts = {"total": len(rows), "with_field_data": sum(bool(r["has_work_data"]) for r in rows),
                  "completion": "Legacy completion status is not standardized."}
        buckets = {str(r["ju"]): "JU_Records" for r in rows}
        readme = ("Project " + job + "\n\nJU_Records contains all current legacy JU notes and photos.\n"
                  "The JU index preserves billing exactly as recorded. Presence of notes, billing or photos "
                  "does not certify completion. Completion is unverified in this legacy packet.\n"
                  "The route workbook is a saved route snapshot; use the JU records for field documentation.\n"
                  "Saved_Reports preserves separately supplied completed-JU and all-JU spreadsheets unchanged; "
                  "those reported legacy statuses have not been independently revalidated by this directory update.\n")
    else:
        records = source / "3_JU_FILES"
        rows = collect(source)
        if not rows:
            raise RuntimeError("No live JUs for " + job)
        if len(rows) != len({str(r["JU"]) for r in rows}):
            raise RuntimeError("Duplicate live JU identifiers")
        for label, selected in (("MASTER", rows),
                ("COMPLETED", [r for r in rows if r["State"] == "COMPLETE"]),
                ("NOT_COMPLETED", [r for r in rows if r["State"] != "COMPLETE"])):
            _write(sheet / (job + "_" + label + ".xlsx"), selected)
        billing_workbook(source, rows, sheet / (job + "_BILLING.xlsx"))
        buckets = {str(r["JU"]): "JU_Completed" if r["State"] == "COMPLETE" else "JU_Not_Completed" for r in rows}
        counts = {"total": len(rows), "completed": sum(r["State"] == "COMPLETE" for r in rows)}
        counts["not_completed"] = counts["total"] - counts["completed"]
        readme = ("Project " + job + "\n\nSpreadsheets: original assignment, current master, completed, not completed, and route.\n"
                  "JU_Completed: closed JUs meeting the saved closing-task photo requirements; includes qualifying surveys/trip charges.\n"
                  "JU_Not_Completed: pending or incomplete JUs.\nEach JU folder contains saved notes, pole details and original photos.\n")
    for p in sheet.iterdir():
        mapping["Spreadsheets/" + p.name] = p
    seen = set()
    for folder in sorted(records.iterdir()):
        if not folder.is_dir():
            continue
        ju = folder.name.split(" - ", 1)[0]
        if ju not in buckets:
            raise RuntimeError("JU missing from index: " + folder.name)
        seen.add(ju)
        for p in sorted(folder.rglob("*")):
            if not p.is_file() or p.is_symlink():
                continue
            rel = p.relative_to(folder)
            if any(x in rel.parts for x in ("VOICE_HISTORY", "RECORD_HISTORY")):
                continue
            allowed = p.suffix.lower() in (IMAGE_TYPES | VIDEO_TYPES) or p.name in (
                "transfer_info.txt", "BILLING_AND_NOTES.txt", "notes and billing codes.txt")
            if allowed:
                mapping[buckets[ju] + "/" + folder.name + "/" + rel.as_posix()] = p
    if seen != set(buckets):
        raise RuntimeError("JU directory/index mismatch")
    (stage / "README.txt").write_text(readme)
    mapping["README.txt"] = stage / "README.txt"
    counts["photos"] = sum(p.suffix.lower() in IMAGE_TYPES for p in mapping.values())
    return mapping, counts

def make_parts(job, files, stage, previous):
    """Each PART_N.zip opens independently; each is below 100 MB."""
    groups, batch, size = [], [], 0
    for relative, path in sorted(files.items()):
        estimated = path.stat().st_size + len(relative.encode()) * 2 + 400
        if estimated > PART_LIMIT:
            raise RuntimeError("Individual file exceeds ZIP part limit: " + str(path))
        if batch and size + estimated > PART_LIMIT:
            groups.append(batch)
            batch, size = [], 0
        batch.append((relative, path))
        size += estimated
    if batch:
        groups.append(batch)
    results = {}
    output = DIRECTORY / "Send_Ready" / job
    output.mkdir(parents=True, exist_ok=True)
    for index, group in enumerate(groups, 1):
        name = "PART_" + str(index) + ".zip"
        signature = hashlib.sha256(json.dumps([(r, sha(p)) for r, p in group]).encode()).hexdigest()
        target = output / name
        old = previous.get(name, {})
        if old.get("signature") == signature and target.exists() and sha(target) == old.get("sha256"):
            results[name] = old
            continue
        tmp = stage / name
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_STORED, allowZip64=True) as z:
            for rel, path in group:
                z.write(path, arcname=job + "/" + rel)
        if tmp.stat().st_size >= 100_000_000:
            raise RuntimeError("ZIP part limit exceeded")
        with zipfile.ZipFile(tmp) as z:
            if z.testzip() is not None:
                raise RuntimeError("ZIP validation failed")
        # These ZIPs are derived from already backed-up canonical files.
        # Preserve manually changed archives, but do not version every regenerated ZIP.
        if target.exists() and old.get("sha256") != sha(target):
            remember(target)
        digest = sha(tmp)
        staged = target.with_name("." + target.name + ".tmp")
        shutil.copyfile(tmp, staged)
        if sha(staged) != digest:
            raise RuntimeError("ZIP transfer validation failed")
        staged.replace(target)
        results[name] = {"signature": signature, "sha256": digest, "bytes": target.stat().st_size}
    for old_name, meta in previous.items():
        target = output / old_name
        if old_name not in results and target.exists():
            if sha(target) != meta.get("sha256"):
                remember(target)
            target.unlink()
    atomic_text(output / "README.txt", "Project " + job + "\nSend all " + str(len(results)) +
                " PART files together. Each ZIP opens independently and is under 100 MB.\n"
                "These filenames are reused when the project is updated.\n")
    return results

def update_all(force=False):
    REPO.mkdir(exist_ok=True)
    with (REPO / ".devco_directory.lock").open("a") as guard:
        try:
            fcntl.flock(guard, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            return {"busy": True}
        state = json.loads(STATE.read_text()) if STATE.exists() else {}
        jobs = sorted(p.name for p in (REPO / "JOBS").iterdir() if p.is_dir()
                      and ((p / "3_JU_FILES").is_dir() or (p / "JOB_PACKET_CURRENT").is_dir()))
        summary = {}
        for job in jobs:
            before = fingerprint(job)
            old = state.get(job, {})
            if not force and old.get("fingerprint") == before:
                summary[job] = old.get("counts", {})
                continue
            with TemporaryDirectory(prefix="devco-publish-") as tmp:
                stage = Path(tmp)
                with RECORD_LOCK:
                    mapping, counts = build_mapping(job, stage)
                    target_root = DIRECTORY / "Jobs" / job
                    generated = {}
                    for rel, source in mapping.items():
                        generated[rel] = publish_file(source, target_root / rel)
                if fingerprint(job) != before:
                    raise RuntimeError("Source changed during export; retrying on next refresh")
                # Remove only files generated by the previous successful run.
                for rel in old.get("files", {}):
                    target = target_root / rel
                    if rel not in generated and target.exists():
                        remember(target)
                        target.unlink()
                for folder in sorted([p for p in target_root.rglob("*") if p.is_dir()],
                                     key=lambda p: len(p.parts), reverse=True):
                    try:
                        folder.rmdir()
                    except OSError:
                        pass
                current = {rel: target_root / rel for rel in generated}
                parts = make_parts(job, current, stage, old.get("parts", {}))
                state[job] = {"fingerprint": before, "files": generated, "parts": parts,
                              "counts": counts, "updated": datetime.now().isoformat(timespec="seconds")}
                atomic_text(STATE, json.dumps(state, indent=2))
                summary[job] = counts
                print("Project directory updated", job, json.dumps(counts), flush=True)
        atomic_text(DIRECTORY / "UPDATE_STATUS.json", json.dumps(
            {j: {"updated": state[j]["updated"], **state[j]["counts"]} for j in jobs if j in state}, indent=2))
        return summary

def watch():
    while True:
        try:
            from gps_photo_audit import refresh as refresh_gps_audit
            refresh_gps_audit()
        except Exception as error:
            print("GPS PHOTO AUDIT ERROR:", error, flush=True)
        try:
            update_all()
        except Exception as error:
            print("PROJECT DIRECTORY UPDATE ERROR:", error, flush=True)
        time.sleep(30)

def start_directory_watcher():
    thread = threading.Thread(target=watch, daemon=True, name="devco-project-directory")
    thread.start()
    return thread

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--watch", action="store_true")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    if args.watch:
        watch()
    else:
        print(json.dumps(update_all(force=args.force), indent=2))
