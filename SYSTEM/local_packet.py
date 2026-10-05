"""Read the live job packet on this phone without uploading or copying records."""
import html
from datetime import datetime
from pathlib import Path
from urllib.parse import urlencode

from job_records import _parse_info, _parse_record
from note_store import read_note
from record_editor import RECORD_LOCK, folder_for, photo_items


def esc(value):
    return html.escape(str(value if value is not None else ''), quote=True)


def url(route, **params):
    return route + '?' + urlencode(params)


STYLE = '''*{box-sizing:border-box}body{margin:0;background:#09131b;color:#edf4f7;font:16px system-ui;line-height:1.5}main{max-width:1000px;margin:auto;padding:18px 16px 50px}header{display:flex;gap:10px;align-items:center;justify-content:space-between}header strong{color:#4cdd92}h1{font-size:28px;line-height:1.2;margin:22px 0 8px}h2{font-size:20px;margin:24px 0 12px}p{overflow-wrap:anywhere}small,.muted{color:#adc1cd}.muted{font-size:14px}a,button{color:#edf4f7}a{text-decoration:none}nav,.row{display:flex;gap:8px;flex-wrap:wrap}.button,button{display:inline-block;padding:11px 15px;min-height:46px;border-radius:10px;background:#203846;border:1px solid #456071;text-align:center;font:600 15px system-ui;cursor:pointer}.primary,button[aria-pressed=true]{background:#4cdd92;color:#06180e;border-color:#4cdd92}.stats{display:grid;grid-template-columns:repeat(3,1fr);gap:8px;margin:20px 0}.stats div,article,section{background:#13232e;border:1px solid #304754;border-radius:12px;padding:16px}.stats b{display:block;font-size:26px}.stats span{font-size:13px;color:#adc1cd}.green{color:#4cdd92}.amber{color:#ffc26c}input{width:100%;padding:14px;background:#0e1d26;color:#edf4f7;border:1px solid #456071;border-radius:10px;font:16px system-ui;margin:14px 0}article{margin:10px 0}article a{display:block}article h3{font-size:20px;margin:0 0 4px}article p{margin:4px 0}.badge{font-size:12px;font-weight:700}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(240px,1fr));gap:12px}.grid a{display:block;padding:15px;border:1px solid #456071;border-radius:10px;background:#13232e}.text{white-space:pre-wrap;overflow-wrap:anywhere}.gallery{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,280px),1fr));gap:14px}figure{margin:0;padding:10px;background:#13232e;border-radius:12px}img{width:100%;height:auto;display:block;border-radius:8px}figcaption{font-size:13px;color:#adc1cd;padding-top:7px;overflow-wrap:anywhere}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:9px;border:1px solid #456071;text-align:left;vertical-align:top;min-width:100px;max-width:400px;white-space:pre-wrap;overflow-wrap:anywhere}th{background:#203846}td a{color:#71d9ff;text-decoration:underline}.tablewrap{overflow:auto;margin-top:15px}details{margin-top:18px}summary{cursor:pointer;padding:10px 0}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:14px system-ui}#empty[hidden],article[hidden]{display:none}@media(max-width:420px){main{padding:14px 12px}h1{font-size:25px}.stats div{padding:12px 8px}.stats b{font-size:23px}}'''


def shell(title, body, job):
    return ('<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            '<title>' + esc(title) + ' | DEVCO</title><style>' + STYLE + '</style></head><body><main>'
            '<header><strong>DEVCO FIELD</strong><span class="muted">Local job packet</span></header>'
            '<nav style="margin-top:16px"><a class="button" href="' + esc(url('/packet', job=job)) + '">Job packet</a>'
            '<a class="button" href="/map">Field map</a><a class="button" href="/">Field</a></nav>' + body + '</main></body></html>')


def job_root(app, job):
    if job not in app['jobs']():
        raise ValueError('Job not found')
    root = app['JOBS'] / job
    if root.resolve().parent != app['JOBS'].resolve():
        raise ValueError('Job not found')
    return root


def spreadsheet_paths(root):
    return sorted(p for section in ('0_ORIGINAL', '1_JOB_WORKFLOW', '2_ROUTE')
                  for p in (root / section).glob('*.xlsx')
                  if p.is_file() and not p.is_symlink())


def spreadsheet_file(app, job, relative):
    root = job_root(app, job)
    path = root / relative
    if path not in spreadsheet_paths(root) or path.resolve().parent.parent != root.resolve():
        raise ValueError('Spreadsheet not found')
    return path


def packet_page(app, job):
    root = job_root(app, job)
    rows = []
    with RECORD_LOCK:
        for info in sorted((root / '3_JU_FILES').glob('*/transfer_info.txt')):
            data = _parse_info(info)
            ju = data.get('JU Record', info.parent.name.split(' - ', 1)[0])
            record = _parse_record(info.parent / 'BILLING_AND_NOTES.txt')
            rows.append({'ju': ju, 'address': data.get('Address', ''), 'state': app['_ju_state'](info.parent),
                         'close': record['status'], 'note': read_note(info.parent),
                         'photos': len(photo_items(info.parent, job, ju))})
    done = sum(r['state'] == 'COMPLETE' for r in rows)
    photos = sum(r['photos'] for r in rows)
    sheets = []
    labels = {'SOURCE': 'Original assignment', 'MASTER': 'All JUs', 'COMPLETED': 'Completed JUs',
              'NOT_COMPLETED': 'Not completed JUs', 'WORKFLOW': 'Job workflow', 'ROUTE': 'Route sheet'}
    for p in spreadsheet_paths(root):
        label = labels.get(p.stem.removeprefix(job + '_'), p.stem)
        sheets.append('<a href="' + esc(url('/packet-sheet', job=job, file=p.relative_to(root).as_posix())) + '"><b>' + esc(label) + '</b><br><small>Open spreadsheet</small></a>')
    cards = []
    for r in rows:
        complete = r['state'] == 'COMPLETE'
        label = 'Completed / closed' if complete else ('Pending' if r['state'] == 'PENDING' else 'Not completed')
        search = ' '.join(str(r[k]) for k in ('ju', 'address', 'close', 'note')).lower()
        cards.append('<article data-state="' + ('complete' if complete else 'open') + '" data-search="' + esc(search) + '">'
                     '<a href="' + esc(url('/packet-ju', job=job, ju=r['ju'])) + '"><span class="badge ' + ('green' if complete else 'amber') + '">' + label + '</span>'
                     '<h3>JU ' + esc(r['ju']) + '</h3><p>' + esc(r['address']) + '</p><p class="muted">' + esc(r['close'] or 'No saved closing task') + '</p>'
                     '<p class="muted">' + str(r['photos']) + ' photos · Open notes &amp; files</p></a></article>')
    body = '<h1>Job ' + esc(job) + '</h1><p class="muted">Current files saved on this phone. Reopen or refresh to see new work.</p>'
    body += '<div class="stats"><div><b>' + str(len(rows)) + '</b><span>Total JUs</span></div><div><b class="green">' + str(done) + '</b><span>Completed / closed</span></div><div><b class="amber">' + str(len(rows)-done) + '</b><span>Not completed</span></div></div>'
    body += '<p>' + str(photos) + ' photos in JU folders</p><h2>Spreadsheets</h2><div class="grid">' + ''.join(sheets) + '</div>'
    body += '<h2>JU files</h2><div class="row" role="group" aria-label="Filter JUs"><button data-filter="all" aria-pressed="true">All (' + str(len(rows)) + ')</button><button data-filter="complete" aria-pressed="false">Completed (' + str(done) + ')</button><button data-filter="open" aria-pressed="false">Not completed (' + str(len(rows)-done) + ')</button></div>'
    body += '<label for="search" class="muted">Find a JU, address, or note</label><input type="search" id="search" placeholder="Search this job" autocomplete="off"><p id="result-count" class="muted" aria-live="polite"></p><div id="records">' + ''.join(cards) + '</div><p id="empty" hidden>No matching JUs.</p>'
    body += '''<script>let filter='all';const cards=[...document.querySelectorAll('article[data-state]')],search=document.getElementById('search');function apply(){const query=search.value.trim().toLowerCase();let count=0;cards.forEach(card=>{card.hidden=!((filter==='all'||card.dataset.state===filter)&&card.dataset.search.includes(query));if(!card.hidden)count++;});document.getElementById('result-count').textContent=count+' JUs shown';document.getElementById('empty').hidden=count!==0;}document.querySelectorAll('[data-filter]').forEach(button=>button.addEventListener('click',()=>{filter=button.dataset.filter;document.querySelectorAll('[data-filter]').forEach(b=>b.setAttribute('aria-pressed',String(b===button)));apply();}));search.addEventListener('input',apply);apply();</script>'''
    return shell('Job ' + job, body, job)


def ju_page(app, job, ju):
    job_root(app, job)
    with RECORD_LOCK:
        folder = folder_for(app, job, ju)
        info = _parse_info(folder / 'transfer_info.txt')
        record = _parse_record(folder / 'BILLING_AND_NOTES.txt')
        note = read_note(folder)
        photos = photo_items(folder, job, ju)
        state = app['_ju_state'](folder)
        raw = (folder / 'transfer_info.txt').read_text(errors='replace')
        record_file = folder / 'BILLING_AND_NOTES.txt'
        saved = record_file.read_text(errors='replace') if record_file.exists() else ''
    label = 'Completed / closed' if state == 'COMPLETE' else ('Pending' if state == 'PENDING' else 'Not completed')
    body = '<h1>JU ' + esc(ju) + '</h1><p>' + esc(info.get('Address', '')) + '</p><p class="muted">Job ' + esc(job) + '</p>'
    body += '<section><b class="' + ('green' if state == 'COMPLETE' else 'amber') + '">' + label + '</b><p>' + esc(record['status'] or 'No saved closing task') + '</p>'
    body += '<h2>Billing codes</h2><p>' + (esc(' · '.join(c + ' ×' + q for c, q in record['billing'])) or 'No saved billing codes') + '</p><h2>Notes</h2><div class="text">' + esc(note or 'No saved notes') + '</div></section>'
    body += '<div class="row" style="margin-top:16px"><a class="button" href="' + esc(url('/record', job=job, ju=ju)) + '">Edit this JU</a>'
    try:
        lat, lon = float(info.get('Latitude', '')), float(info.get('Longitude', ''))
        if -90 <= lat <= 90 and -180 <= lon <= 180:
            body += '<a class="button" href="' + esc('https://www.google.com/maps/dir/?' + urlencode({'api': 1, 'destination': f'{lat},{lon}', 'travelmode': 'driving'})) + '">Navigate to pole</a>'
    except ValueError:
        pass
    body += '</div><h2>Photos (' + str(len(photos)) + ')</h2><div class="gallery">'
    for p in photos:
        body += '<figure><a href="' + esc(p['url']) + '"><img loading="lazy" decoding="async" src="' + esc(p['url']) + '" alt="JU ' + esc(ju) + ': ' + esc(p['name']) + '"></a><figcaption>' + esc(p['name']) + '</figcaption></figure>'
    body += '</div>' + ('<p class="muted">No photos saved in this JU folder.</p>' if not photos else '<p class="muted">Tap a photo to open the full-size original.</p>')
    body += '<details><summary>Original pole details</summary><pre>' + esc(raw) + '</pre></details>'
    if saved:
        body += '<details><summary>Saved billing and notes file</summary><pre>' + esc(saved) + '</pre></details>'
    return shell('JU ' + ju, body, job)


def sheet_page(app, job, relative, sheet_name=''):
    from openpyxl import load_workbook
    path = spreadsheet_file(app, job, relative)
    with RECORD_LOCK:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            name = sheet_name or wb.sheetnames[0]
            if name not in wb.sheetnames:
                raise ValueError('Worksheet not found')
            sheet = wb[name]
            tabs = ''.join('<a class="button' + (' primary' if n == name else '') + '" href="' + esc(url('/packet-sheet', job=job, file=relative, sheet=n)) + '">' + esc(n) + '</a>' for n in wb.sheetnames)
            table = []
            for index, values in enumerate(sheet.iter_rows(values_only=True)):
                tag = 'th' if index == 0 else 'td'
                cells = []
                for value in values:
                    rendered = esc(value)
                    if isinstance(value, str) and value.startswith(('https://', 'http://')):
                        rendered = '<a href="' + esc(value) + '">Open link</a>'
                    cells.append('<' + tag + '>' + rendered + '</' + tag + '>')
                table.append('<tr>' + ''.join(cells) + '</tr>')
        finally:
            wb.close()
    modified = datetime.fromtimestamp(path.stat().st_mtime).astimezone().strftime('%b %d, %Y at %I:%M %p %Z')
    body = '<h1>' + esc(path.name) + '</h1><p class="muted">Saved spreadsheet preview · Updated ' + esc(modified) + '</p><p class="muted">JU files show the latest saved field notes and photos.</p>'
    body += '<nav>' + tabs + '<a class="button" href="' + esc(url('/packet-file', job=job, file=relative)) + '">Download Excel file</a></nav>'
    body += '<p class="muted">Swipe sideways to see all columns. This preview displays saved cell values.</p><div class="tablewrap"><table>' + ''.join(table) + '</table></div>'
    return shell(path.name, body, job)
