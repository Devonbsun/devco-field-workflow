"""Regression tests use disposable job folders, never production records."""
import tempfile, unittest, threading, urllib.request, urllib.parse, json
from pathlib import Path
from unittest.mock import patch
import devco_app as app

class SimpleFlow(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        app.JOBS=Path(self.tmp.name)/"JOBS"
        app.STATE=Path(self.tmp.name)/"state.json"
        app.active_job="TEST";app.active_ju="101"
        app.photo_notice="Watching for new Solocator photos"
        self.folder=app.JOBS/"TEST"/"3_JU_FILES"/"101 - Test"
        self.folder.mkdir(parents=True)
        (self.folder/"transfer_info.txt").write_text("JU Record: 101\nAddress: Test street\nLatitude: 41.6\nLongitude: -93.6\n")
        second=self.folder.parent/"102 - Next";second.mkdir()
        (second/"transfer_info.txt").write_text("JU Record: 102\nAddress: Next street\nLatitude: 41.7\nLongitude: -93.7\n")
        (self.folder/"photos").mkdir()
        (self.folder/"photos"/"one.jpg").write_bytes(b"test")
        (self.folder/"photos"/"two.jpg").write_bytes(b"test")
        self.server=app.ThreadingHTTPServer(("127.0.0.1",0),app.H)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.base="http://127.0.0.1:"+str(self.server.server_port)
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.tmp.cleanup()
    def get(self,path):
        return urllib.request.urlopen(self.base+path).read().decode()
    def test_pages_status_and_form_identity(self):
        home=self.get("/")
        self.assertIn("1. Take photos",home)
        self.assertNotIn("FIELD TOOLS",home)
        self.assertIn("watchPosition",home)
        self.assertEqual(json.loads(self.get("/status"))["photos"],2)
        billing=self.get("/billing")
        self.assertIn('name="job" value="TEST"',billing)
        self.assertIn('name="ju" value="101"',billing)
        self.assertIn("localStorage",billing)
        self.assertIn("leaflet",self.get("/map"))
    def test_gps_cannot_replace_selected_pole(self):
        state=json.loads(self.get("/gps-nearest?lat=41.7&lon=-93.7&accuracy=5"))
        self.assertFalse(state["switched"]);self.assertEqual(app.active_ju,"101")
    def test_stale_form_does_not_save(self):
        data=urllib.parse.urlencode(dict(job="TEST",ju="102",close="TRANSFER ALREADY COMPLETED")).encode()
        response=urllib.request.urlopen(self.base+"/finish",data=data).read().decode()
        self.assertIn("Nothing was saved",response)
        self.assertFalse((self.folder/"BILLING_AND_NOTES.txt").exists())
    def test_successful_closeout_and_next_pole(self):
        with patch.object(app,"sync_job") as sync:
            data=urllib.parse.urlencode(dict(job="TEST",ju="101",close="FIBER TRANSFER COMPLETED",code="WC1F",qty_WC1F="1",note="test")).encode()
            response=urllib.request.urlopen(self.base+"/finish",data=data)
            self.assertIn("saved=TEST:101",response.url)
            self.assertEqual(app.active_ju,"102")
            self.assertIn("WC1F x1",(self.folder/"BILLING_AND_NOTES.txt").read_text())
            sync.assert_called_once()
    def test_photo_matching_and_validation(self):
        photo=Path(self.tmp.name)/"new.jpg";photo.write_bytes(b"image")
        with patch.object(app,"_photo_gps",return_value=(41.6,-93.6)),patch.object(app,"sync_job"):
            result=app.import_solocator_photo(photo,quiet=True)
            self.assertEqual(result["ju"],"101")
            self.assertTrue((self.folder/"photos"/"new.jpg").exists())
        self.assertFalse(app.save_closeout("TEST","101","PENDING",[],"")[0])
        self.assertFalse(app.save_closeout("TEST","101","FIBER TRANSFER COMPLETED",[],"")[0])

    def test_auto_location_selects_job_and_ju(self):
        app.active_job=None;app.active_ju=None
        result=json.loads(self.get("/gps-nearest?auto=1&lat=41.7&lon=-93.7&accuracy=5"))
        self.assertTrue(result["switched"])
        self.assertEqual((app.active_job,app.active_ju),("TEST","102"))
    def test_auto_location_rejects_poor_and_far_gps(self):
        for query in ["lat=41.7&lon=-93.7&accuracy=100","lat=40&lon=-90&accuracy=5"]:
            result=json.loads(self.get("/gps-nearest?auto=1&"+query))
            self.assertFalse(result["switched"])
            self.assertEqual(app.active_ju,"101")
    def test_auto_location_rejects_ambiguous_jobs(self):
        import shutil
        shutil.copytree(app.JOBS/"TEST",app.JOBS/"OTHER")
        result=json.loads(self.get("/gps-nearest?auto=1&lat=41.7&lon=-93.7&accuracy=5"))
        self.assertFalse(result["switched"])
        self.assertIn("Two poles",result["message"])
    def test_auto_location_keeps_next_stop_after_completion(self):
        with patch.object(app,"sync_job"):
            app.save_closeout("TEST","101","FIBER TRANSFER COMPLETED",[("WC1F","1")],"")
        app.active_ju="102"
        result=json.loads(self.get("/gps-nearest?auto=1&lat=41.6&lon=-93.6&accuracy=5"))
        self.assertFalse(result["switched"])
        self.assertEqual(app.active_ju,"102")

    def test_note_autosave_survives_reopen_and_active_change(self):
        note="Existing transfer.\nCustomer confirmed access."
        data=urllib.parse.urlencode(dict(job="TEST",ju="101",note=note)).encode()
        self.assertTrue(json.loads(urllib.request.urlopen(self.base+"/note",data=data).read())["ok"])
        self.assertIn(note,self.get("/billing"))
        app.active_ju="102"
        data=urllib.parse.urlencode(dict(job="TEST",ju="101",note="Updated old JU note")).encode()
        urllib.request.urlopen(self.base+"/note",data=data).read()
        from note_store import read_note
        self.assertEqual(read_note(self.folder),"Updated old JU note")
        self.assertEqual(app.active_ju,"102")
    def test_already_completed_preserves_full_note_in_export(self):
        note="Transfer was already completed.\nNo work performed by Devco."
        data=urllib.parse.urlencode(dict(job="TEST",ju="101",close="TRANSFER ALREADY COMPLETED",note=note)).encode()
        urllib.request.urlopen(self.base+"/finish",data=data).read()
        from job_records import _parse_record
        rec=_parse_record(self.folder/"BILLING_AND_NOTES.txt")
        self.assertEqual(rec["notes"],[note])
        self.assertEqual(rec["billing"],[])
        from openpyxl import load_workbook
        wb=load_workbook(app.JOBS/"TEST"/"1_JOB_WORKFLOW"/"TEST_MASTER.xlsx")
        self.assertEqual(wb.active.cell(2,10).value,note);wb.close()
        app.active_ju="101"
        self.assertIn(note,self.get("/billing"))
    def test_failed_closeout_keeps_note(self):
        for p in (self.folder/"photos").glob("*"):p.unlink()
        data=urllib.parse.urlencode(dict(job="TEST",ju="101",close="TRANSFER ALREADY COMPLETED",note="Keep this note")).encode()
        from urllib.error import HTTPError
        with self.assertRaises(HTTPError):urllib.request.urlopen(self.base+"/finish",data=data)
        from note_store import read_note
        self.assertEqual(read_note(self.folder),"Keep this note")
        self.assertFalse((self.folder/"BILLING_AND_NOTES.txt").exists())

    def test_adss_and_no_services_keep_selected_billing(self):
        from job_records import _parse_record, collect
        (self.folder/"photos"/"two.jpg").unlink()
        for close in ("ADSS", "NO SERVICES ON POLE"):
            app.active_ju="101"
            data=urllib.parse.urlencode(dict(job="TEST",ju="101",close=close,code="WC1",qty_WC1="9",note="Survey documented")).encode()
            urllib.request.urlopen(self.base+"/finish",data=data).read()
            record=_parse_record(self.folder/"BILLING_AND_NOTES.txt")
            self.assertEqual(record["status"],close)
            self.assertEqual(record["billing"],[("WC1","9")])
            self.assertEqual(record["notes"],["Survey documented"])
            row=next(r for r in collect(app.JOBS/"TEST") if r["JU"]=="101")
            self.assertEqual(row["State"],"COMPLETE")
            self.assertEqual(row["Billing Codes"],"WC1 x9")
        (self.folder/"photos"/"one.jpg").unlink()
        for close in ("ADSS", "NO SERVICES ON POLE"):
            self.assertFalse(app.save_closeout("TEST","101",close,[],"")[0])

if __name__=="__main__":unittest.main()
