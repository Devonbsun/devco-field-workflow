"""Completed JU exports grouped by completion or original photo day."""
from collections import Counter, defaultdict
from datetime import datetime
import html
import io
import json
from pathlib import Path
import re
import shutil
import tempfile
import threading
import uuid
import zipfile
from central_time import CENTRAL

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from job_records import CLOSE_RULES, _parse_record
from record_editor import RECORD_LOCK, atomic_text
import tracking as t
import tracking_export as packets

UNKNOWN = 'DATE_NEEDED'
LOCK = threading.Lock()
ACTIVE = set()
LIMIT = 10_000_000


def completion_day(value):
    if not value:
        return UNKNOWN
    try:
        stamp = datetime.fromisoformat(value.strip().replace('Z', '+00:00'))
        # Older app timestamps were local and had no offset.
        stamp = stamp.replace(tzinfo=CENTRAL) if stamp.tzinfo is None else stamp.astimezone(CENTRAL)
        return stamp.date().isoformat()
    except (ValueError, TypeError):
        return UNKNOWN


def photo_date(path):
    """Original capture time only; a copied file's mtime is not a work date."""
    from PIL import Image
    try:
        with Image.open(path) as image:
            exif = image.getexif()
            details = exif.get_ifd(34665)
            value = details.get(36867) or exif.get(36867)
            if value:
                stamp = datetime.strptime(str(value), '%Y:%m:%d %H:%M:%S')
                offset = details.get(36881) or exif.get(36881)
                return completion_day(stamp.isoformat() + (str(offset) if offset else '')), 'Original photo timestamp'
    except (OSError, ValueError, KeyError, TypeError):
        pass
    match = re.search(r'Solocator[-_](20\d{2}-\d{2}-\d{2})[ T_]', path.name, re.I)
    if match:
        day = completion_day(match.group(1))
        if day != UNKNOWN: return day, 'Solocator photo filename timestamp'
    return UNKNOWN, 'Date needed'


def records(root, job):
    rows, legacy_count = [], 0
    with RECORD_LOCK:
        for row in t.records(root, job):
            if row['legacy']:
                legacy_count += 1
                continue
            if CLOSE_RULES.get(row['close_code'], (None, None))[1] != 'COMPLETE' or not row['checks']['Photos']:
                continue
            source = row['_folder'] / 'BILLING_AND_NOTES.txt'
            row['completed_at'] = _parse_record(source)['completed_at']
            row['day'] = completion_day(row['completed_at'])
            row['date_source'] = 'Saved completion timestamp'
            if row['day'] == UNKNOWN:
                dates = [photo_date(p) for p in row['_photos']]
                dates = [(d, s) for d, s in dates if d != UNKNOWN]
                row['day'], row['date_source'] = max(dates, default=(UNKNOWN, 'Date needed'))
            row['_record_hash'] = packets.digest(source)
            rows.append(row)
    return rows, legacy_count


def status_path(root, job):
    if job not in t.jobs(root):
        raise ValueError('Job not found.')
    return Path(root) / 'TRACKING/daily-exports' / (job + '.json')


def save(root, job, value):
    file = status_path(root, job)
    file.parent.mkdir(parents=True, exist_ok=True)
    atomic_text(file, json.dumps(value, indent=2))


def state(root, job):
    rows, legacy = records(root, job)
    file = status_path(root, job)
    saved = json.loads(file.read_text()) if file.exists() else None
    if saved and saved['status'] == 'building' and job not in ACTIVE:
        saved.update(status='failed', error='App restarted during export. Tap Export again.')
        save(root, job, saved)
    counts = Counter(r['day'] for r in rows)
    return {'job': job, 'jobs': t.jobs(root), 'completed': len(rows), 'legacy': legacy,
            'days': [{'day': day, 'count': counts[day]} for day in sorted(counts)], 'export': saved}


def workbook(headers, values, title):
    wb = Workbook(); ws = wb.active; ws.title = title
    ws.append(headers)
    for row in values:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Field text must never become an Excel formula.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='163A4B')
    for col, label in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = 48 if label in ('Address', 'Notes', 'Photo files') else 28 if label in ('Close code', 'Billing codes') else 21
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime): cell.number_format = 'mm/dd/yy'
        if row[0].row % 2 == 0:
            for cell in row: cell.fill = PatternFill('solid', fgColor='EDF5F7')
        ws.row_dimensions[row[0].row].height = 60
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO(); wb.save(out); return out.getvalue()


def date_value(day):
    return 'Date needed' if day == UNKNOWN else datetime.fromisoformat(day)


def billing_rows(rows):
    return [[r['job'], r['ju'], r['address'], date_value(r['day']), r['date_source'], line['code'], int(line['quantity']) if line['quantity'].isdigit() else line['quantity']]
            for r in rows for line in r['lines']]


BILL_HEADERS = ['Project', 'JU', 'Address', 'Work date', 'Date source', 'Billing code', 'Quantity']


def report(row):
    text = [f"PROJECT: {row['job']}", f"JU: {row['ju']}", f"ADDRESS: {row['address']}",
            'WORK DATE: ' + ('Date needed' if row['day'] == UNKNOWN else row['day']),
            'DATE SOURCE: ' + row['date_source'],
            f"CLOSE CODE: {row['close_code']}", f"COORDINATES: {row['latitude']}, {row['longitude']}",
            '', 'BILLING CODES:']
    text += [f"{line['code']} x{line['quantity']}" for line in row['lines']] or ['No saved billing codes']
    text += ['', 'SAVED NOTES:', row['note'] or 'No saved notes', '', f"PHOTOS: {row['photo_count']}"]
    return ('\n'.join(text) + '\n').encode()


def folder_name(row):
    address = re.sub(r'[^A-Za-z0-9 .,_-]', '_', row['address']).strip(' .')[:90]
    return 'JU_' + row['ju'] + (' - ' + address if address else '')


def destination(root, job):
    return packets.share_root(root) / job / 'Completed_By_Day'


def create(root, job):
    with LOCK:
        if ACTIVE:
            raise t.Conflict('An export is already being prepared.')
        rows, _ = records(root, job)
        if not rows:
            raise ValueError('No verified completed JUs in this project. Older unverified records remain in the job packet.')
        ident = uuid.uuid4().hex
        value = {'id': ident, 'status': 'building', 'count': len(rows), 'progress': 0}
        ACTIVE.add(job)
        try: save(root, job, value)
        except Exception:
            ACTIVE.discard(job)
            raise
        threading.Thread(target=build, args=(Path(root), job, rows, value), daemon=True, name='daily-JU-export').start()
        return {'ok': True}


def build(root, job, rows, value, limit=LIMIT):
    stage = None
    try:
        with packets.WORKER:
            dest = destination(root, job); dest.parent.mkdir(parents=True, exist_ok=True)
            stage = Path(tempfile.mkdtemp(prefix='.daily-stage-', dir=dest.parent))
            grouped = defaultdict(list)
            for row in rows: grouped[row['day']].append(row)
            days, source_hashes = [], {}
            for day, day_rows in sorted(grouped.items()):
                daily = stage / day; daily.mkdir()
                photo_manifest = []
                for row in day_rows:
                    ju_folder = daily / '2_JU_Files' / folder_name(row)
                    photos = ju_folder / 'Photos'; photos.mkdir(parents=True)
                    (ju_folder / 'JU_REPORT.txt').write_bytes(report(row))
                    # Only explicitly saved field records appear in boss packets.
                    for n, source in enumerate(row['_photos'], 1):
                        compact, checksum = packets.compact_photo(root, source)
                        source_hashes[source] = checksum
                        name = f'{n:03d}_' + re.sub(r'[^A-Za-z0-9 .,_-]', '_', source.stem)[:100] + compact.suffix.lower()
                        target = photos / name; shutil.copy2(compact, target)
                        photo_manifest.append({'ju': row['ju'], 'file': target.relative_to(daily).as_posix(),
                                               'original_name': source.name, 'original_sha256': checksum,
                                               'shared_sha256': packets.digest(target)})
                    bill = daily / '3_Billing_Only' / (folder_name(row) + '.csv'); bill.parent.mkdir(exist_ok=True)
                    data = [[v.isoformat()[:10] if isinstance(v, datetime) else v for v in line] for line in billing_rows([row])]
                    bill.write_bytes(packets.csv_bytes([BILL_HEADERS] + data))
                    value['progress'] += 1; save(root, job, value)
                sheets = daily / '1_Spreadsheets'; sheets.mkdir()
                data = [[r['job'], r['ju'], r['address'], date_value(day), r['date_source'], 'Completed', r['close_code'], r['photo_count'],
                         '; '.join(l['code'] + ' x' + l['quantity'] for l in r['lines']), r['note'],
                         '; '.join(p['file'] for p in photo_manifest if p['ju'] == r['ju'])] for r in day_rows]
                (sheets / 'Completed_JUs.xlsx').write_bytes(workbook(['Project', 'JU', 'Address', 'Work date', 'Date source', 'Status', 'Close code', 'Photo count', 'Billing codes', 'Notes', 'Photo files'], data, 'Completed JUs'))
                (sheets / 'Billing_Codes.xlsx').write_bytes(workbook(BILL_HEADERS, billing_rows(day_rows), 'Billing codes'))
                (daily / 'PHOTO_MANIFEST.json').write_text(json.dumps(photo_manifest, indent=2))
                (daily / 'README.txt').write_text(f'Project {job}\nWork day: {day}\n{len(day_rows)} completed JUs\n\n1_Spreadsheets: completed JUs and billing codes.\n2_JU_Files: one folder per JU with saved notes, close code, billing codes and all photos.\n3_Billing_Only: one billing-only file per JU.\n\nDates use saved completion timestamps first, otherwise the latest original photo timestamp (or original Solocator filename). Date source is shown in each report and spreadsheet. Timestamps without an offset are treated as local Central time. DATE_NEEDED means no valid timestamp is available. File modification dates are never used.\nPhotos are reduced copies; full-resolution originals remain in field records.\nIf this day has several ZIP parts, extract every part into the same folder. Spreadsheets and billing files occur once.\n')
                # Split only at file boundaries; every archive opens independently.
                files = sorted(f for f in daily.rglob('*') if f.is_file())
                groups, group, size = [], [], 1000
                for file in files:
                    estimate = file.stat().st_size + 1000
                    if estimate >= limit: raise ValueError('A file exceeds the packet limit: ' + file.name)
                    if group and size + estimate >= limit:
                        groups.append(group); group = []; size = 1000
                    group.append(file); size += estimate
                if group: groups.append(group)
                parts, coverage = [], []
                for n, group in enumerate(groups, 1):
                    name = f'{job}_{day}_PART_{n}.zip'; target = daily / name
                    with zipfile.ZipFile(target, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                        for file in group:
                            rel = file.relative_to(daily).as_posix()
                            archive.write(file, day + '/' + rel); coverage.append(rel)
                    with zipfile.ZipFile(target) as archive:
                        if archive.testzip(): raise ValueError('ZIP integrity check failed.')
                    if target.stat().st_size >= limit: raise ValueError('Packet size check failed.')
                    parts.append({'name': name, 'day': day, 'bytes': target.stat().st_size, 'sha256': packets.digest(target)})
                if sorted(coverage) != sorted(f.relative_to(daily).as_posix() for f in files):
                    raise ValueError('Export coverage check failed.')
                days.append({'day': day, 'count': len(day_rows), 'photos': len(photo_manifest), 'parts': parts})
            # A change to a completion date or record invalidates the export too.
            current, _ = records(root, job)
            signature = lambda rs: sorted((r['ju'], r['revision'], r['_record_hash']) for r in rs)
            if signature(current) != signature(rows) or any(packets.digest(p) != h for p, h in source_hashes.items()):
                raise t.Conflict('Field records changed during export. Tap Export again for the latest records.')
            backup = (Path('/storage/emulated/0/DEVCO/Backups/Daily_Exports') if packets.phone_root(root) else Path(root) / 'TRACKING/daily-history')
            old = backup / (job + '-' + value['id'])
            if dest.exists():
                backup.mkdir(parents=True, exist_ok=True); shutil.move(str(dest), old)
            try: stage.rename(dest)
            except Exception:
                if old.exists(): shutil.move(str(old), dest)
                raise
            stage = None
            value.update(status='ready', days=days, folder=str(dest), created=t.now())
            save(root, job, value)
    except Exception as error:
        value.update(status='failed', error=str(error)); save(root, job, value)
    finally:
        if stage: shutil.rmtree(stage, ignore_errors=True)
        ACTIVE.discard(job)


def packet_file(root, job, ident, name):
    file = status_path(root, job)
    saved = json.loads(file.read_text()) if file.exists() else {}
    if saved.get('status') != 'ready' or ident != saved.get('id'):
        raise ValueError('Open the latest daily export.')
    part = next((p for d in saved['days'] for p in d['parts'] if p['name'] == name), None)
    if not part or Path(name).name != name: raise ValueError('File not found.')
    path = destination(root, job) / part['day'] / name
    if not path.is_file() or packets.digest(path) != part['sha256']:
        raise ValueError('Export file changed or missing. Tap Export again.')
    return path
